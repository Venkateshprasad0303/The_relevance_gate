"""Generate the pitch deck (relevance_gate_deck.pptx) with speaker notes.
Run: py make_deck.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PAPER = RGBColor(0xF5, 0xF2, 0xEA)
INK = RGBColor(0x18, 0x15, 0x11)
SOFT = RGBColor(0x5B, 0x55, 0x4A)
CLEAR = RGBColor(0x2E, 0x6B, 0x4E)
HOLD = RGBColor(0xA0, 0x3A, 0x26)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(kicker, title, bullets, notes, accent=INK):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER

    if kicker:
        kb = s.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.5))
        p = kb.text_frame.paragraphs[0]
        r = p.add_run(); r.text = kicker.upper()
        r.font.name = "Consolas"; r.font.size = Pt(14); r.font.color.rgb = SOFT

    tb = s.shapes.add_textbox(Inches(0.85), Inches(1.05), Inches(11.6), Inches(2.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Georgia"; r.font.size = Pt(40); r.font.bold = False
    r.font.color.rgb = accent

    if bullets:
        bb = s.shapes.add_textbox(Inches(0.9), Inches(3.2), Inches(11.5), Inches(3.8))
        bf = bb.text_frame; bf.word_wrap = True
        for i, (lead, rest) in enumerate(bullets):
            p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            p.space_after = Pt(16)
            r = p.add_run(); r.text = "—  "
            r.font.name = "Georgia"; r.font.size = Pt(20); r.font.color.rgb = SOFT
            r = p.add_run(); r.text = lead
            r.font.name = "Georgia"; r.font.size = Pt(20); r.font.bold = True
            r.font.color.rgb = INK
            if rest:
                r = p.add_run(); r.text = "  " + rest
                r.font.name = "Georgia"; r.font.size = Pt(20); r.font.color.rgb = SOFT

    s.notes_slide.notes_text_frame.text = notes
    return s


# 1 — title
s = slide(
    "The Relevance Gate",
    "Job-hunt outreach that refuses to send unless it finds a real reason.",
    [("Found many. Contact few.", "All of them true.")],
    "SAY (10s): Every outreach tool in this room will send anything to anyone. "
    "I built the one that says no. The Relevance Gate is an agent crew that "
    "refuses to send a message unless it finds a real, evidence-backed reason "
    "to contact someone.",
)

# 2 — problem
slide(
    "The problem",
    "Cold outreach is dead on arrival.",
    [
        ("~3.4% reply rate", "for generic cold outreach."),
        ("0.1 seconds", "is how long a recruiter looks before deleting — ~90% discarded on sight."),
        ("15–25% reply rate", "when outreach is signal-based and evidence-grounded. A ~5× lift."),
        ("The volume death-loop", "torches your sender reputation and your reputation, period."),
    ],
    "SAY (20s): Generic cold outreach replies at about three and a half percent. "
    "Recipients decide in a tenth of a second and ninety percent dies on sight. "
    "But evidence-grounded, signal-based outreach replies at fifteen to "
    "twenty-five percent — five X. Every tool today chases volume and fakes "
    "personalization. That's the loop that kills job seekers' credibility.",
)

# 3 — the idea
slide(
    "The idea",
    "We optimize for send-worthiness. The headline feature is the message we won't send.",
    [
        ("The gate can say no.", "If there's no genuine overlap, nothing is written. On purpose."),
        ("The message you don't send", "protects the ones you do."),
    ],
    "SAY (15s): So we flipped the objective. Instead of optimizing volume, we "
    "optimize send-worthiness. Our headline feature is the message we won't "
    "send. No other agent here declines to act — ours treats refusal as the "
    "product.",
    accent=HOLD,
)

# 4 — architecture
slide(
    "How it works",
    "Four isolated agents. The gate rules before the writer exists.",
    [
        ("RESEARCH", "builds a sourced dossier from what the target said, their page, and opt-in live web search."),
        ("INTERSECT", "computes the genuine overlap between your proven work and their world — evidenced both ways."),
        ("THE GATE", "an independent verifier: scores 0–100, sends or HOLDS. Fails closed."),
        ("WRITE", "only if cleared: ≤75 words, one low-friction ask, every claim traceable to evidence."),
    ],
    "SAY (25s): Four isolated Claude agents. Research builds a sourced dossier. "
    "Intersect finds the genuine overlap between my proven work and their "
    "world. Then the gate — an independent verifier — scores it and decides. "
    "The key design choice: the gate renders its verdict BEFORE the writer "
    "exists, so the writer can never talk the gate into sending. That "
    "independence is the anti-hallucination property. And in code, the writer "
    "physically never runs on a weak overlap. If anything fails to parse, it "
    "fails CLOSED — an unverifiable message does not send.",
)

# 5 — live demo
slide(
    "Live demo",
    "Tell me your name and what you work on.",
    [
        ("A real judge, live:", "their words become the evidence. Watch the crew find the overlap and clear the gate."),
        ("Then the refusal:", "a backend engineer who 'thinks AI is neat' — held, score 0. Generic interest is not a reason."),
    ],
    "SAY (60s — THE DEMO): [Turn to a judge] What's your name, and what do you "
    "work on? [Type it in. Hit Run.] Live: research, then the intersection — "
    "there's the overlap: they work on agents, I shipped a 7-phase multi-agent "
    "legal reasoning pipeline with 6 anti-hallucination layers. The gate "
    "clears it at 85, and the message opens on that exact overlap — "
    "sixty-some words, one ask. [Read the first line aloud.] "
    "Now the moment: [type the bank engineer who merely 'thinks AI is neat'] "
    "— the gate HOLDS. Score zero. No message. On purpose. It found no real "
    "reason, so it sent nothing.",
    accent=CLEAR,
)

# 6 — the full loop
slide(
    "The full system",
    "Résumé in → real people found → researched, scored, and only then written to.",
    [
        ("Paste your résumé or a URL", "— an agent distills it into your profile with strict relevance criteria."),
        ("Find targets", "— live web search surfaces people publicly hiring or building in your specialty, each with a source."),
        ("One click per candidate", "— full research, an honest 0–100 score, and a perfect message or a principled refusal."),
    ],
    "SAY (25s): And it's a complete system. Paste your résumé — it becomes "
    "your profile. Type who you're looking for — live web search finds real "
    "people with real sources: hiring posts, founders building in your "
    "specialty. One click runs any of them through the gate. Found many, "
    "contacted few, all of them true. Works for anyone's career — zero code "
    "change.",
)

# 7 — engineering
slide(
    "Engineered for the stage",
    "Reliability is the feature.",
    [
        ("Verdict in ~10 seconds;", "schema-enforced JSON on every call — no parse roulette."),
        ("Fails closed:", "no verdict means no send. No raw error can ever reach the screen."),
        ("Every verdict logged to ClickHouse", "— block-rate by overlap strength, queryable live."),
        ("Demo safety net:", "if the API dies mid-pitch, a pre-validated result still renders."),
    ],
    "SAY (20s): Under the hood: every agent call is schema-enforced, the "
    "verdict lands in about ten seconds, failures degrade gracefully, and the "
    "system fails closed — which is the only safe direction for an anti-spam "
    "agent. Every verdict is logged to ClickHouse, so 'it refuses' isn't a "
    "claim — it's a queryable block-rate.",
)

# 8 — sponsors + vision
slide(
    "The stack & the vision",
    "An agent allowed to say no.",
    [
        ("Anthropic", "— Claude is the brain of all four agents."),
        ("ClickHouse", "— logs why it refuses, at scale."),
        ("Render", "— shipped live, deployed from this repo."),
        ("Guild", "— the whole entry is one idea: an agent allowed to say no."),
        ("The vision:", "a relevance layer any outreach tool can plug into. Built the wedge; this is the layer."),
    ],
    "SAY (20s): Claude is the brain of every agent. ClickHouse logs why it "
    "refuses at scale. It ships live on Render. And the Guild story is the "
    "whole entry: an agent allowed to say no. Today it's the job hunt — the "
    "wedge I needed myself. The company is a relevance layer any outreach "
    "tool can plug into. Found many. Contacted few. All of them true. Thank you.",
    accent=CLEAR,
)

prs.save("relevance_gate_deck.pptx")
print("Saved relevance_gate_deck.pptx —", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
